# -*- coding: utf-8 -*-
"""Status prodaje iz tablice na Google Driveu -> status.json (samo oznake i status, bez kupaca).
ID tablice je u tajni SHEET_ID (GitHub secret), nikad u kodu stranice.  SHEET_ID=... python sync_status.py"""
import io
import json
import os
import sys
import urllib.request

from openpyxl import load_workbook

sid = os.environ["SHEET_ID"].strip()
url = "https://docs.google.com/spreadsheets/d/%s/export?format=xlsx" % sid
data = urllib.request.urlopen(urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"}), timeout=60).read()
wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
out = {"prodano": [], "rezervirano": []}
for ime in ("Stanovi", "Garaža i ostave"):
    ws = wb[ime]
    rows = ws.iter_rows(values_only=True)
    zag = [str(c or "").strip().lower() for c in next(rows)]
    io_, is_ = zag.index("oznaka"), zag.index("status")
    for r in rows:
        if not r or len(r) <= max(io_, is_):          # Google skraćuje prazne retke
            continue
        oz, st = r[io_], str(r[is_] or "").strip().lower()
        if not oz:
            continue
        if st.startswith("prodan"):
            out["prodano"].append(str(oz).strip())
        elif st.startswith("rezerv"):
            out["rezervirano"].append(str(oz).strip())
if len(out["prodano"]) + len(out["rezervirano"]) > 102:
    sys.exit("previše oznaka - krivi list?")
tekst = json.dumps(out, ensure_ascii=False, indent=2) + "\n"
put = sys.argv[1] if len(sys.argv) > 1 else "status.json"
staro = open(put, encoding="utf-8").read() if os.path.exists(put) else ""
if staro != tekst:
    open(put, "w", encoding="utf-8").write(tekst)
    print("promjena:", out)
else:
    print("bez promjene")
